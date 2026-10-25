"""Build trimmed real fixtures for the loaders (runs in GitHub Actions, which has network access).

Picks German ETS installations of a few known groups plus some others, and follows them through
the EEA Industrial Reporting database (ETS id -> installation -> facility -> parent company).
Also prints link statistics for all German installations so docs can quote measured numbers.
"""

from __future__ import annotations

import csv
import io
import json
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path
from urllib.parse import urljoin

import httpx
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).parents[1] / "src"))
from belegt.resolution import normalise_lei, normalise_register

UA = {"User-Agent": "belegt-portfolio/0.1 (+https://github.com/Parsaka/belegt)"}
OUT = Path("inspect/fixtures")
OUT.mkdir(parents=True, exist_ok=True)
client = httpx.Client(headers=UA, timeout=300, follow_redirects=True)
GROUPS = re.compile(r"BASF|Heidelberg Materials|HeidelbergCement|RWE", re.I)
UR_PAGE = "https://climate.ec.europa.eu/areas-action/carbon-markets/eu-emissions-trading-system-eu-ets/union-registry_en"
ACCDB = sys.argv[1] if len(sys.argv) > 1 else "/tmp/eea.accdb"


def ets_rows() -> tuple[list[tuple[object, ...]], int, list[tuple[object, ...]]]:
    html = client.get(UR_PAGE).text
    links = re.findall(r'href="([^"]+)"', html)
    url = next(urljoin(UR_PAGE, x) for x in links if x.endswith("verified_emissions_2025_en.xlsx"))
    wb = load_workbook(io.BytesIO(client.get(url).content), read_only=True, data_only=True)
    for ws in wb.worksheets:  # the data sheet is not the first one
        rows = list(ws.iter_rows(values_only=True))
        hdr_i = next((i for i, r in enumerate(rows) if r and r[0] == "REGISTRY_CODE"), None)
        if hdr_i is not None:
            print(f"ETS sheet {ws.title!r} of {wb.sheetnames}")
            return rows[:hdr_i], hdr_i, rows[hdr_i:]
    raise SystemExit("no sheet with a REGISTRY_CODE header")


def ets_key(installation_id: object) -> str:
    """Union Registry id 1231 -> EEA ETSIdentifier DE000000000001231."""
    return f"DE{int(str(installation_id).strip()):015d}"


def export(table: str) -> list[dict[str, str]]:
    raw = subprocess.run(["mdb-export", ACCDB, table], capture_output=True, check=True).stdout
    return list(csv.DictReader(io.StringIO(raw.decode("utf-8", "replace"))))


def write_csv(name: str, rows: list[dict[str, str]]) -> None:
    if not rows:
        print("empty:", name)
        return
    with (OUT / name).open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def main() -> None:
    preamble, hdr_i, table = ets_rows()
    hdr = table[0]
    col = {name: i for i, name in enumerate(hdr)}
    de = [
        r
        for r in table[1:]
        if r and r[0] == "DE" and str(r[col["ACCOUNT_TYPE"]]).upper().startswith("OPERATOR HOLDING")
    ]
    if not de:  # account type labels differ between releases
        de = [r for r in table[1:] if r and r[0] == "DE" and not r[col["IMO_NUMBER"]]]
    ets_ids = export("3f2_ETSIdentifiers")
    eea_de = {r["ETSIdentifier"] for r in ets_ids if r["ETSIdentifier"].startswith("DE")}
    # all installations the EEA links (Brandenburg only, see link_stats), a few known group
    # sites by name, and an even spread of the rest
    picked = [r for r in de if ets_key(r[col["INSTALLATION_IDENTIFIER"]]) in eea_de]
    op_header, op_rows = operators_table()
    op_col = {name: k for k, name in enumerate(op_header)}
    group_ids = {
        str(o[op_col["Installation ID"]])
        for o in op_rows
        if GROUPS.search(str(o[op_col["Account Holder Name"]]))
    }
    picked += [
        r for r in de if r not in picked and str(r[col["INSTALLATION_IDENTIFIER"]]) in group_ids
    ][:15]
    rest = [r for r in de if r not in picked]
    picked += rest[:: max(1, len(rest) // 10)][:10]
    print(f"ETS: header at row {hdr_i}, {len(de)} German stationary rows, {len(picked)} picked")

    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    for r in [*preamble, hdr, *picked]:
        ws.append(list(r))
    wb.save(OUT / "ets_verified_emissions_de.xlsx")

    inst = {r["Installation_INSPIRE_ID"]: r for r in export("3_ProductionInstallation")}
    fac = {r["Facility_INSPIRE_ID"]: r for r in export("2_ProductionFacility")}
    de_ets = [r for r in ets_ids if r["ETSIdentifier"].startswith("DE")]

    # link statistics over all German ETS installations (latest EEA year per id)
    latest: dict[str, dict[str, str]] = {}
    for r in de_ets:
        if (
            r["ETSIdentifier"] not in latest
            or r["reportingYear"] > latest[r["ETSIdentifier"]]["reportingYear"]
        ):
            latest[r["ETSIdentifier"]] = r
    ur_ids = {ets_key(r[col["INSTALLATION_IDENTIFIER"]]) for r in de}
    linked = [latest[i] for i in ur_ids if i in latest]
    with_parent = [
        r
        for r in linked
        if (
            f := fac.get(
                inst.get(r["Installation_INSPIRE_ID"], {}).get("Parent_Facility_INSPIRE_ID", "")
            )
        )
        and f.get("parentCompanyName")
    ]
    parents = Counter(
        fac[inst[r["Installation_INSPIRE_ID"]]["Parent_Facility_INSPIRE_ID"]]["parentCompanyName"]
        for r in with_parent
    )
    stats = {
        "ets_de_stationary_installations": len(ur_ids),
        "eea_de_ets_identifiers": len(latest),
        "ets_installations_found_in_eea": len(linked),
        "with_parent_company_name": len(with_parent),
        "distinct_parent_company_names": len(parents),
        "eea_reporting_years": sorted({r["reportingYear"] for r in de_ets}),
        "top_parents": parents.most_common(25),
    }
    print(json.dumps(stats, indent=1, ensure_ascii=False))
    (OUT / "link_stats.json").write_text(json.dumps(stats, indent=1, ensure_ascii=False))

    want = {ets_key(r[col["INSTALLATION_IDENTIFIER"]]) for r in picked}
    fx_ets = [r for r in de_ets if r["ETSIdentifier"] in want]
    fx_inst = [
        inst[i] for i in dict.fromkeys(r["Installation_INSPIRE_ID"] for r in fx_ets) if i in inst
    ]
    fx_fac = [
        fac[f] for f in dict.fromkeys(r["Parent_Facility_INSPIRE_ID"] for r in fx_inst) if f in fac
    ]
    write_csv("eea_ets_identifiers.csv", fx_ets)
    write_csv("eea_installations.csv", fx_inst)
    write_csv("eea_facilities.csv", fx_fac)
    print(
        f"fixtures: {len(fx_ets)} ETS ids, {len(fx_inst)} installations, {len(fx_fac)} facilities"
    )
    for f in fx_fac:
        print(
            "   ",
            f["Facility_INSPIRE_ID"],
            "|",
            f["nameOfFeature"],
            "|",
            f["parentCompanyName"],
            "|",
            f["city"],
        )
    picked_ids = {str(r[col["INSTALLATION_IDENTIFIER"]]) for r in picked}
    fx_ops = [o for o in op_rows if str(o[op_col["Installation ID"]]) in picked_ids]
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    for o in [op_header, *fx_ops]:
        ws.append(list(o))
    wb.save(OUT / "ets_operators_de.xlsx")
    print(f"operators fixture: {len(fx_ops)} rows")
    leis = sorted({lei for o in fx_ops if (lei := normalise_lei(o[op_col["LEI"]]))})
    regs = sorted(
        {
            reg
            for o in fx_ops
            if (reg := normalise_register(o[op_col["Company Registration Nr of Account Holder"]]))
        }
    )
    names = {f["parentCompanyName"] for f in fx_fac if f.get("parentCompanyName")}
    names |= {str(o[op_col["Account Holder Name"]]) for o in fx_ops if not o[op_col["LEI"]]}
    gleif(sorted(names), leis, regs)


OPERATORS_FILE = "policy_ets_registry_operators_ets_en.xlsx"


def operators_table() -> tuple[tuple[object, ...], list[tuple[object, ...]]]:
    """The Commission's list of installation operators: account holder, register number, LEI."""
    html = client.get(UR_PAGE).text
    url = next(
        urljoin(UR_PAGE, x)
        for x in re.findall(r'href="([^"]+)"', html)
        if x.endswith(OPERATORS_FILE)
    )
    wb = load_workbook(io.BytesIO(client.get(url).content), read_only=True, data_only=True)
    print("operators workbook:", url, wb.sheetnames)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    header = rows[0]
    print("operators header:", header)
    col = {name: k for k, name in enumerate(header)}
    de = [r for r in rows[1:] if r and r[col["MS Registry"]] == "DE"]
    with_lei = [r for r in de if r[col["LEI"]]]
    with_reg = [r for r in de if r[col["Company Registration Nr of Account Holder"]]]
    regs = Counter(
        re.sub(r"\s*\d.*", "", str(r[col["Company Registration Nr of Account Holder"]]))
        for r in with_reg
    )
    stats = {
        "operators_rows_total": len(rows) - 1,
        "operators_de": len(de),
        "operators_de_with_lei": len(with_lei),
        "operators_de_with_registration_nr": len(with_reg),
        "operators_de_registration_prefixes": regs.most_common(10),
        "operators_de_distinct_holders": len({r[col["Account Holder Name"]] for r in de}),
    }
    print(json.dumps(stats, indent=1, ensure_ascii=False))
    (OUT / "operator_stats.json").write_text(json.dumps(stats, indent=1, ensure_ascii=False))
    for r in de[:5]:
        print("   ", r)
    return header, de


GLEIF = "https://api.gleif.org/api/v1"


def gleif(parent_names: list[str], leis: list[str], regs: list[str]) -> None:
    """LEI records for the operators (by LEI, by register number, by name) and the real RR rows
    under three groups."""
    candidates: dict[str, list[dict[str, object]]] = {}
    for i in range(0, len(leis), 100):
        r = client.get(
            f"{GLEIF}/lei-records",
            params={"filter[lei]": ",".join(leis[i : i + 100]), "page[size]": 100},
        )
        candidates[f"lei:{i}"] = r.json().get("data", [])
    print(
        f"   GLEIF records by operator LEI: {sum(len(v) for v in candidates.values())} of {len(leis)}"
    )
    for reg in regs[:80]:
        found: list[dict[str, object]] = []
        # GLEIF stores the number as registered, e.g. "HRB 6000" or "HRB 19220 P"
        variants = {reg, re.sub(r"(\d)([A-Z])$", r"\1 \2", reg)}
        for variant in variants:
            r = client.get(
                f"{GLEIF}/lei-records",
                params={
                    "filter[entity.registeredAs]": variant,
                    "filter[entity.legalAddress.country]": "DE",
                    "page[size]": 10,
                },
            )
            found += r.json().get("data", []) if r.status_code == 200 else []
        candidates[f"reg:{reg}"] = found
        print(
            f"   GLEIF by register number {reg!r}: {[c['attributes']['entity']['legalName']['name'] for c in candidates[f'reg:{reg}']]}"
        )
    for name in parent_names:
        r = client.get(f"{GLEIF}/lei-records", params={"filter[fulltext]": name, "page[size]": 5})
        candidates[name] = r.json().get("data", []) if r.status_code == 200 else []
        print(
            f"   GLEIF candidates for {name!r}: {[c['attributes']['entity']['legalName']['name'] for c in candidates[name]]}"
        )
    (OUT / "gleif_candidates.json").write_text(json.dumps(candidates, ensure_ascii=False))

    groups = {}
    for name in ("BASF SE", "Heidelberg Materials AG", "RWE Aktiengesellschaft"):
        r = client.get(
            f"{GLEIF}/lei-records", params={"filter[entity.legalName]": name, "page[size]": 1}
        )
        if data := r.json().get("data"):
            groups[data[0]["id"]] = name
    print("   groups:", groups)

    gc = client.get("https://goldencopy.gleif.org/api/v2/golden-copies/publishes/latest").json()[
        "data"
    ]
    import zipfile

    rr_rows: list[dict[str, str]] = []
    blob = io.BytesIO(client.get(gc["rr"]["full_file"]["csv"]["url"]).content)
    with zipfile.ZipFile(blob) as z, z.open(z.namelist()[0]) as fh:
        reader = csv.DictReader(io.TextIOWrapper(fh, "utf-8"))
        for row in reader:
            ends = (row["Relationship.EndNode.NodeID"], row["Relationship.StartNode.NodeID"])
            if any(lei in groups for lei in ends):
                rr_rows.append(row)
    keep = [c for c in reader.fieldnames or [] if not re.search(r"Qualifiers|Quantifiers", c)]
    rr_rows = [{k: r[k] for k in keep} for r in rr_rows]
    print(
        f"   RR rows touching the groups: {len(rr_rows)}; types {Counter(r['Relationship.RelationshipType'] for r in rr_rows)}"
    )
    write_csv("gleif_rr_groups.csv", rr_rows)

    leis = sorted({r["Relationship.StartNode.NodeID"] for r in rr_rows} | set(groups))
    records: list[dict[str, object]] = []
    for i in range(0, len(leis), 100):
        r = client.get(
            f"{GLEIF}/lei-records",
            params={"filter[lei]": ",".join(leis[i : i + 100]), "page[size]": 100},
        )
        records += r.json().get("data", [])
    (OUT / "gleif_entities_groups.json").write_text(json.dumps(records, ensure_ascii=False))
    print(f"   GLEIF entity records saved: {len(records)}")


if __name__ == "__main__":
    main()
