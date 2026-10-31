"""Trim the GeoNames German postal-code file to the postcodes in the operators fixture.

Runs in GitHub Actions (download.geonames.org is not reachable from the development sandbox).
Keeps the original tab-separated lines unchanged and prints how many German ETS operator
postcodes the full file covers, so docs can quote a measured number.
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path
from urllib.parse import urljoin

import httpx
from openpyxl import load_workbook

UA = {"User-Agent": "belegt-portfolio/0.1 (+https://github.com/Parsaka/belegt)"}
GEONAMES = "https://download.geonames.org/export/zip/DE.zip"
UR_PAGE = "https://climate.ec.europa.eu/areas-action/carbon-markets/eu-emissions-trading-system-eu-ets/union-registry_en"
OPERATORS_FILE = "policy_ets_registry_operators_ets_en.xlsx"
FIXTURE = Path("tests/fixtures/ets_operators_de.xlsx")
OUT = Path("inspect/fixtures")


def postcodes(rows: list[tuple[object, ...]], registry: str | None = None) -> set[str]:
    col = {str(name).strip(): i for i, name in enumerate(rows[0])}
    return {
        str(r[col["Contact PCode"]]).strip().zfill(5)
        for r in rows[1:]
        if r
        and r[col["Contact PCode"]]
        and (registry is None or r[col["MS Registry"]] == registry)
    }


def sheet(blob: bytes | Path) -> list[tuple[object, ...]]:
    src = io.BytesIO(blob) if isinstance(blob, bytes) else blob
    wb = load_workbook(src, read_only=True, data_only=True)
    return list(wb.worksheets[0].iter_rows(values_only=True))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    client = httpx.Client(headers=UA, timeout=300, follow_redirects=True)
    with zipfile.ZipFile(io.BytesIO(client.get(GEONAMES).content)) as z:
        lines = z.read("DE.txt").decode("utf-8").splitlines()
    by_code: dict[str, list[str]] = {}
    for line in lines:
        by_code.setdefault(line.split("\t")[1], []).append(line)

    wanted = postcodes(sheet(FIXTURE))
    kept = [line for code in sorted(wanted) for line in by_code.get(code, [])]
    (OUT / "geonames_postcodes_de.txt").write_text("\n".join(kept) + "\n", encoding="utf-8")

    html = client.get(UR_PAGE).text
    url = next(
        urljoin(UR_PAGE, x)
        for x in re.findall(r'href="([^"]+)"', html)
        if x.endswith(OPERATORS_FILE)
    )
    all_de = postcodes(sheet(client.get(url).content), registry="DE")
    stats = {
        "geonames_lines": len(lines),
        "geonames_distinct_postcodes": len(by_code),
        "fixture_postcodes": len(wanted),
        "fixture_lines_kept": len(kept),
        "operators_de_distinct_postcodes": len(all_de),
        "operators_de_postcodes_in_geonames": len(all_de & set(by_code)),
        "operators_de_postcodes_missing": sorted(all_de - set(by_code))[:30],
        "first_lines": lines[:3],
    }
    print(json.dumps(stats, indent=1, ensure_ascii=False))
    (OUT / "geonames_stats.json").write_text(json.dumps(stats, indent=1, ensure_ascii=False))


if __name__ == "__main__":
    main()
