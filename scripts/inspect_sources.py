"""Discover and describe the raw source files (run in GitHub Actions, which has network access).

Prints what each publisher actually serves: file links, sheet names, header rows, row counts.
Writes trimmed samples to `inspect/` so they can become test fixtures.
"""

from __future__ import annotations

import io
import json
import re
import sys
import zipfile
from pathlib import Path
from urllib.parse import urljoin

import httpx
from openpyxl import load_workbook

UA = {"User-Agent": "belegt-portfolio/0.1 (+https://github.com/Parsaka/belegt)"}
OUT = Path("inspect")
OUT.mkdir(exist_ok=True)
client = httpx.Client(headers=UA, timeout=120, follow_redirects=True)


def section(title: str) -> None:
    print(f"\n{'=' * 20} {title} {'=' * 20}", flush=True)


def describe_xlsx(name: str, content: bytes) -> None:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        print(f"-- sheet {ws.title!r}: {len(rows)} rows")
        for i, row in enumerate(rows[:25]):
            print(f"   {i:>3}: {row}")
        sample = [list(r) for r in rows[:60]]
        (OUT / f"{name}__{ws.title}.json").write_text(json.dumps(sample, default=str))


def union_registry() -> None:
    section("EU ETS Union Registry")
    page = "https://climate.ec.europa.eu/areas-action/carbon-markets/eu-emissions-trading-system-eu-ets/union-registry_en"
    html = client.get(page).text
    links = sorted(set(re.findall(r'href="([^"]+\.(?:xlsx|xls)[^"]*)"', html)))
    links += sorted(set(re.findall(r'href="([^"]*/document/download/[^"]+)"', html)))
    for link in links:
        print("link:", urljoin(page, link))
    for link in links:
        url = urljoin(page, link)
        if "2025" in url or "2024" in url or len(links) < 6:
            r = client.get(url)
            print(f"\nGET {url} -> {r.status_code} {r.headers.get('content-type')} {len(r.content)} bytes")
            if r.status_code == 200 and r.content[:2] == b"PK":
                try:
                    describe_xlsx(Path(url.split("?")[0]).stem[:40], r.content)
                except Exception as exc:  # discovery script: report and continue
                    print("  could not parse:", exc)


def union_registry_site() -> None:
    section("union-registry-data.ec.europa.eu")
    for url in ["https://union-registry-data.ec.europa.eu/", "https://union-registry-data.ec.europa.eu/report/welcome"]:
        r = client.get(url)
        print(url, r.status_code, r.headers.get("content-type"))
        print(r.text[:1500])


def eea() -> None:
    section("EEA Industrial Reporting")
    base = "https://sdi.eea.europa.eu/webdav/datastore/public/"
    html = client.get(base).text
    dirs = sorted(set(re.findall(r'href="[^"]*?(eea_t_ied-eprtr[^"/]*)/?"', html)))
    print("datasets:", dirs)
    if not dirs:
        print(html[:2000])
        return

    def version(d: str) -> int:
        m = re.search(r"_v(\d+)_", d)
        return int(m.group(1)) if m else -1

    latest = max(dirs, key=version)
    print("latest:", latest)
    listing = client.get(base + latest + "/").text
    files = re.findall(r'href="([^"?]+)"', listing)
    for f in files:
        print("  file:", f)
    for f in files:
        if f.lower().endswith(".zip") and ("csv" in f.lower() or "user" in f.lower()):
            url = urljoin(base + latest + "/", f)
            r = client.get(url)
            print(f"GET {url} -> {r.status_code} {len(r.content)} bytes")
            with zipfile.ZipFile(io.BytesIO(r.content)) as z:
                for info in z.infolist():
                    print(f"   {info.filename} {info.file_size}")
                    if info.filename.lower().endswith(".csv"):
                        with z.open(info) as fh:
                            head = fh.read(4000).decode("utf-8", "replace").splitlines()[:6]
                        for line in head:
                            print("      ", line[:400])
                        (OUT / Path(info.filename).name).write_text("\n".join(head))


def gleif() -> None:
    section("GLEIF")
    r = client.get(
        "https://api.gleif.org/api/v1/lei-records",
        params={"filter[entity.legalName]": "BASF SE", "page[size]": 1},
    )
    print("lei search", r.status_code)
    data = r.json()
    print(json.dumps(data, indent=1)[:3000])
    if data.get("data"):
        lei = data["data"][0]["id"]
        for rel in ("direct-parent", "ultimate-parent", "direct-children"):
            rr = client.get(f"https://api.gleif.org/api/v1/lei-records/{lei}/{rel}", params={"page[size]": 3})
            print(rel, rr.status_code, rr.text[:800])
    gc = client.get("https://goldencopy.gleif.org/api/v2/golden-copies/publishes/latest")
    print("golden copy latest", gc.status_code)
    print(gc.text[:3000])


if __name__ == "__main__":
    steps = {"ur": union_registry, "ursite": union_registry_site, "eea": eea, "gleif": gleif}
    for key in sys.argv[1:] or steps:
        try:
            steps[key]()
        except Exception as exc:  # discovery script: keep going
            print(f"!! {key} failed: {exc!r}")
