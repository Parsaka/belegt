"""EU ETS Union Registry: verified emissions per installation, wide xlsx to long rows.

The workbook has a few preamble rows, then a header starting with REGISTRY_CODE, then one row per
account with ALLOCATION_<year>, VERIFIED_EMISSIONS_<year> and CH_* (Swiss linking) columns.
Missing values are text such as "n/a"; they are kept as a status, never as zero.
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

_YEAR_COL = re.compile(r"^(?P<kind>ALLOCATION|VERIFIED_EMISSIONS)_?(?P<year>\d{4})$")


@dataclass(frozen=True)
class EtsEmission:
    registry_code: str
    installation_id: str
    account_holder: str
    installation_name: str
    permit_id: str | None
    activity_code: int | None
    account_type: str | None
    year: int
    verified_t: int | None
    allocation_t: int | None
    status: str  # "reported", or the publisher's text such as "n/a"

    @property
    def ets_identifier(self) -> str:
        """The key the EEA database uses: DE + 15-digit zero-padded installation id."""
        return f"{self.registry_code}{int(self.installation_id):015d}"


def _number(value: object) -> tuple[int | None, str]:
    if value is None or value == "":
        return None, "missing"
    if isinstance(value, int | float):
        return int(value), "reported"
    text = str(value).strip()
    if re.fullmatch(r"-?\d+", text):
        return int(text), "reported"
    return None, text.lower()


def _text(value: object) -> str | None:
    text = str(value).strip() if value is not None else ""
    return text or None


def read_rows(rows: list[tuple[object, ...]], registry: str | None = "DE") -> Iterator[EtsEmission]:
    header_at = next(i for i, row in enumerate(rows) if row and row[0] == "REGISTRY_CODE")
    header = [str(h) for h in rows[header_at]]
    col = {name: i for i, name in enumerate(header)}
    years: dict[int, dict[str, int]] = {}
    for i, name in enumerate(header):
        if m := _YEAR_COL.match(name):
            years.setdefault(int(m["year"]), {})[m["kind"]] = i
    for row in rows[header_at + 1 :]:
        if not row or not row[0] or (registry and row[0] != registry):
            continue
        if row[col["IMO_NUMBER"]] not in (None, ""):
            continue  # maritime operator, not an installation
        activity = row[col["MAIN_ACTIVITY_TYPE_CODE"]]
        for year, cols in sorted(years.items()):
            verified, status = (
                _number(row[cols["VERIFIED_EMISSIONS"]])
                if "VERIFIED_EMISSIONS" in cols
                else (None, "missing")
            )
            allocation, _ = (
                _number(row[cols["ALLOCATION"]]) if "ALLOCATION" in cols else (None, "missing")
            )
            yield EtsEmission(
                registry_code=str(row[0]),
                installation_id=str(row[col["INSTALLATION_IDENTIFIER"]]).strip(),
                account_holder=str(row[col["IDENTIFIER_IN_REG"]]).strip(),
                installation_name=str(row[col["INSTALLATION_NAME"]]).strip(),
                permit_id=_text(row[col["PERMIT_IDENTIFIER"]]),
                activity_code=int(activity) if isinstance(activity, int | float) else None,
                account_type=_text(row[col["ACCOUNT_TYPE"]]),
                year=year,
                verified_t=verified,
                allocation_t=allocation,
                status=status,
            )


def read_verified_emissions(path: Path, registry: str | None = "DE") -> Iterator[EtsEmission]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    # The published workbook has explanatory sheets before the data sheet.
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if any(row and row[0] == "REGISTRY_CODE" for row in rows):
            yield from read_rows(rows, registry)
            return
    raise ValueError(f"{path.name}: no sheet with a REGISTRY_CODE header")


@dataclass(frozen=True)
class Operator:
    """One row of the Commission's operator list: who holds the account of an installation."""

    registry_code: str
    installation_id: str
    account_holder: str
    register_nr: str | None  # free text, e.g. "HRB 6000 Amtsgericht Ludwigshafen"
    lei: str | None  # as reported, possibly with dashes
    installation_name: str
    activity_type: str | None
    permit_id: str | None
    permit_status: str | None
    city: str | None
    postcode: str | None


def read_operators(path: Path, registry: str | None = "DE") -> Iterator[Operator]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    col = {str(name).strip(): i for i, name in enumerate(rows[0])}
    for row in rows[1:]:
        if not row or (registry and row[col["MS Registry"]] != registry):
            continue
        yield Operator(
            registry_code=str(row[col["MS Registry"]]),
            installation_id=str(row[col["Installation ID"]]).strip(),
            account_holder=str(row[col["Account Holder Name"]]).strip(),
            register_nr=_text(row[col["Company Registration Nr of Account Holder"]]),
            lei=_text(row[col["LEI"]]),
            installation_name=str(row[col["Installation Name"]]).strip(),
            activity_type=_text(row[col["Activity Type"]]),
            permit_id=_text(row[col["Permit ID"]]),
            permit_status=_text(row[col["Permit Expiry/Revocation Date"]]),
            city=_text(row[col["Contact City"]]),
            postcode=_text(row[col["Contact PCode"]]),
        )
