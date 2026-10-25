from datetime import date
from pathlib import Path

import pytest

from belegt.loaders import eea, ets, gleif

FIX = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="module")
def emissions():
    return list(ets.read_verified_emissions(FIX / "ets_verified_emissions_de.xlsx"))


def test_ets_is_long_format_one_row_per_installation_year(emissions):
    keys = [(e.installation_id, e.year) for e in emissions]
    assert len(keys) == len(set(keys))
    assert {e.year for e in emissions} >= set(range(2008, 2026))


def test_ets_keeps_text_markers_as_status_not_zero(emissions):
    na = [e for e in emissions if e.status == "n/a"]
    assert na and all(e.verified_t is None for e in na)


def test_ets_identifier_matches_eea_format(emissions):
    links = {link.ets_identifier for link in eea.read_ets_links(FIX / "eea_ets_identifiers.csv")}
    ours = {e.ets_identifier for e in emissions}
    assert links and links <= ours


def test_ets_identifier_padding():
    row = ets.EtsEmission("DE", "1231", "", "", None, None, None, 2024, 1, None, "reported")
    assert row.ets_identifier == "DE000000000001231"


def test_ets_skips_maritime_and_other_registries():
    header = ("REGISTRY_CODE", "IDENTIFIER_IN_REG", "INSTALLATION_NAME", "INSTALLATION_IDENTIFIER",
              "PERMIT_IDENTIFIER", "MAIN_ACTIVITY_TYPE_CODE", "VERIFIED_EMISSIONS_2024", "ALLOCATION_2024",
              "IMO_NUMBER", "ACCOUNT_TYPE")  # fmt: skip
    rows = [
        ("preamble",), header,
        ("DE", "a", "Werk", "1", "p", 20, 10, 5, None, "OPERATOR HOLDING ACCOUNT"),
        ("DE", "ship", "Ship", "2", None, 50, 99, None, 9_000_000, "SHIPPING"),
        ("FR", "b", "Usine", "3", "p", 20, 7, 1, None, "OPERATOR HOLDING ACCOUNT"),
    ]  # fmt: skip
    out = list(ets.read_rows(rows))
    assert [(e.installation_id, e.verified_t, e.allocation_t) for e in out] == [("1", 10, 5)]


def test_eea_chain_reaches_parent_company():
    links = list(eea.read_ets_links(FIX / "eea_ets_identifiers.csv"))
    inst = {
        i.installation_inspire_id: i for i in eea.read_installations(FIX / "eea_installations.csv")
    }
    fac = {f.facility_inspire_id: f for f in eea.read_facilities(FIX / "eea_facilities.csv")}
    reached = [
        fac[inst[link.installation_inspire_id].facility_inspire_id]
        for link in links
        if link.installation_inspire_id in inst
    ]
    assert reached and all(f.parent_company_name for f in reached)


def test_gleif_relationships_are_consolidation_only_with_periods():
    rels = list(gleif.read_relationships(FIX / "gleif_rr.csv"))
    assert rels
    assert {r.relationship_type for r in rels} <= gleif.CONSOLIDATION
    assert all(
        r.valid_from is None or r.valid_to is None or r.valid_from <= r.valid_to for r in rels
    )


def test_gleif_period_prefers_relationship_period():
    row = {
        "Relationship.Period.1.periodType": "ACCOUNTING_PERIOD",
        "Relationship.Period.1.startDate": "2024-01-01T00:00:00Z",
        "Relationship.Period.1.endDate": "2024-12-31T00:00:00Z",
        "Relationship.Period.2.periodType": "RELATIONSHIP_PERIOD",
        "Relationship.Period.2.startDate": "2019-07-01T00:00:00Z",
        "Relationship.Period.2.endDate": "",
    }
    assert gleif.period(row) == (date(2019, 7, 1), None)


def test_gleif_entity_from_api_record():
    import json

    records = json.loads((FIX / "gleif_entities_groups.json").read_text())
    basf = next(gleif.entity_from_api(r) for r in records if r["id"] == "529900PM64WH8AF1E917")
    assert basf.legal_name == "BASF SE"
    assert basf.registered_as == "HRB 6000"


def test_verified_emissions_found_behind_explanatory_sheets(tmp_path):
    """The published workbook puts notes before the data sheet; the fixture has only the data."""
    from openpyxl import Workbook, load_workbook

    from belegt.loaders.ets import read_verified_emissions

    source = load_workbook(FIX / "ets_verified_emissions_de.xlsx", read_only=True)
    wb = Workbook()
    wb.active.title = "Notes"
    wb.active.append(["Verified emissions, explanatory notes"])
    data = wb.create_sheet("data")
    for row in source.worksheets[0].iter_rows(values_only=True):
        data.append(list(row))
    wb.save(tmp_path / "full.xlsx")
    assert list(read_verified_emissions(tmp_path / "full.xlsx")) == list(
        read_verified_emissions(FIX / "ets_verified_emissions_de.xlsx")
    )
